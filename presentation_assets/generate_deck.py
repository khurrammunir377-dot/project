import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# 16:9 Widescreen dimensions
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# Color Palette
DARK_BG = RGBColor(11, 17, 32)        # #0B1120
CARD_BG = RGBColor(15, 23, 42)        # #0F172A
BORDER_COLOR = RGBColor(30, 41, 59)   # #1E293B
EMERALD = RGBColor(16, 185, 129)      # #10B981
SKY_BLUE = RGBColor(56, 189, 248)     # #38BDF8
ROYAL_BLUE = RGBColor(37, 99, 235)    # #2563EB
TEXT_WHITE = RGBColor(255, 255, 255)  # #FFFFFF
TEXT_MUTED = RGBColor(148, 163, 184)  # #94A3B8
TEXT_LIGHT = RGBColor(226, 232, 240)  # #E2E8F0

SCREENSHOT_DIR = r"e:\Company App\presentation_assets\screenshots"
OUT_PPTX = r"e:\Company App\Nexora_ERP_Client_Presentation.pptx"

def set_slide_background(slide, prs):
    # Add dark background rectangle
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = DARK_BG
    bg.line.fill.background()
    return bg

def add_header(slide, badge_text, title_text, subtitle_text=""):
    # Header container
    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.3))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    # Badge paragraph
    p0 = tf.paragraphs[0]
    p0.text = badge_text.upper()
    p0.font.size = Pt(9.5)
    p0.font.bold = True
    p0.font.color.rgb = EMERALD
    p0.space_after = Pt(3)

    # Title paragraph
    p1 = tf.add_paragraph()
    p1.text = title_text
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.space_after = Pt(2)

    # Subtitle paragraph
    if subtitle_text:
        p2 = tf.add_paragraph()
        p2.text = subtitle_text
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = TEXT_MUTED

def add_image_with_frame(slide, img_path, left, top, width, height):
    # Background card frame
    frame = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left - Inches(0.08),
        top - Inches(0.08),
        width + Inches(0.16),
        height + Inches(0.16)
    )
    frame.fill.solid()
    frame.fill.fore_color.rgb = CARD_BG
    frame.line.color.rgb = BORDER_COLOR
    frame.line.width = Pt(1.5)

    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, left, top, width, height)

def build_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    blank_layout = prs.slide_layouts[6]

    # ==========================================
    # SLIDE 1: Title Slide
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, prs)

    # Decorative top glow card
    glow = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.73), Inches(5.9))
    glow.fill.solid()
    glow.fill.fore_color.rgb = CARD_BG
    glow.line.color.rgb = BORDER_COLOR
    glow.line.width = Pt(1.5)

    # Title text box
    tb1 = s1.shapes.add_textbox(Inches(1.4), Inches(1.3), Inches(10.5), Inches(4.8))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p_badge = tf1.paragraphs[0]
    p_badge.text = "PAKISTAN'S COMPLETE ENTERPRISE ERP SUITE"
    p_badge.font.size = Pt(11)
    p_badge.font.bold = True
    p_badge.font.color.rgb = EMERALD
    p_badge.space_after = Pt(12)

    p_main = tf1.add_paragraph()
    p_main.text = "Nexora Business Management Platform"
    p_main.font.size = Pt(36)
    p_main.font.bold = True
    p_main.font.color.rgb = TEXT_WHITE
    p_main.space_after = Pt(10)

    p_sub = tf1.add_paragraph()
    p_sub.text = "All-in-one centralized web ERP engineered for Pakistani business environments: FBR Sales Tax Invoicing, Multi-Godown Logistics, EOBI Payroll, and Executive P&L Analytics."
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = TEXT_LIGHT
    p_sub.space_after = Pt(24)

    # Key pillars on title slide
    p_pillars = tf1.add_paragraph()
    p_pillars.text = "✓ FBR Digital POS & Tax Compliance    ✓ Karachi • Lahore • Islamabad Multi-Godown    ✓ Automated EOBI / PESSI Payroll"
    p_pillars.font.size = Pt(11)
    p_pillars.font.bold = True
    p_pillars.font.color.rgb = SKY_BLUE
    p_pillars.space_after = Pt(30)

    p_meta = tf1.add_paragraph()
    p_meta.text = "Live Interactive Demo: https://kmbnexora.vercel.app  •  Confidential Client Presentation"
    p_meta.font.size = Pt(10.5)
    p_meta.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 2: Executive Dashboard
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, prs)
    add_header(s2, "Executive Overview", "Unified Executive Command Center", "Real-time visibility over enterprise cash flow, YTD revenue, pending approvals, and department budgets.")

    # Image (Left)
    add_image_with_frame(s2, os.path.join(SCREENSHOT_DIR, "03_dashboard.png"), Inches(0.8), Inches(1.9), Inches(7.5), Inches(4.9))

    # Text Card (Right)
    tb = s2.shapes.add_textbox(Inches(8.6), Inches(1.9), Inches(3.9), Inches(4.9))
    tf = tb.text_frame
    tf.word_wrap = True
    
    bullets = [
        ("Birds-Eye Financial Visibility", "Realized revenue, monthly operational expenses, and net profit tracked in PKR (Rs.) with live baseline toggles."),
        ("Pending Approvals Radar", "Instant notification of invoices, purchase orders, and expense claims awaiting executive signoff."),
        ("Multi-Metric KPI Cards", "Directly clickable cards allowing directors to dive straight into workforce, stock valuation, or invoices."),
        ("Department Distribution", "Dynamic visual breakdown of staff and expense allocation across Administration, Accounts, and Logistics.")
    ]
    for i, (title, desc) in enumerate(bullets):
        p_t = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p_t.text = f"• {title}"
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = SKY_BLUE
        p_t.space_after = Pt(2)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_LIGHT
        p_d.space_after = Pt(14)

    # ==========================================
    # SLIDE 3: FBR Digital Invoicing & Billing
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, prs)
    add_header(s3, "Finance & Compliance", "FBR-Compliant Digital Invoicing & Tax Ledger", "Automated cryptographic QR codes, provincial sales tax calculations, and instant audit trails.")

    add_image_with_frame(s3, os.path.join(SCREENSHOT_DIR, "04_invoices.png"), Inches(0.8), Inches(1.9), Inches(7.5), Inches(4.9))

    tb = s3.shapes.add_textbox(Inches(8.6), Inches(1.9), Inches(3.9), Inches(4.9))
    tf = tb.text_frame
    tf.word_wrap = True

    bullets = [
        ("FBR Digital POS Integration", "Compliant invoice formatting with automated tax calculation (GST, Provincial SST) ready for Iris reporting."),
        ("Verifiable QR Codes", "Each invoice features a unique cryptographic QR code for authentic digital verification by customers and authorities."),
        ("Accounts Receivable Tracking", "Real-time visibility over Paid, Partially Paid, and Overdue customer invoices with automated reminders."),
        ("1-Click PDF Export & Letterhead", "Generate professional tax invoices on corporate letterhead ready for email and WhatsApp dispatch.")
    ]
    for i, (title, desc) in enumerate(bullets):
        p_t = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p_t.text = f"• {title}"
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = EMERALD
        p_t.space_after = Pt(2)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_LIGHT
        p_d.space_after = Pt(14)

    # ==========================================
    # SLIDE 4: Multi-Item Quotations Generator
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, prs)
    add_header(s4, "Commercial & Sales", "Professional Quotations & Proposal Generator", "Rapid multi-item quote builder with automated tax, discounts, terms, and letterhead.")

    add_image_with_frame(s4, os.path.join(SCREENSHOT_DIR, "05_quotations.png"), Inches(0.8), Inches(1.9), Inches(7.5), Inches(4.9))

    tb = s4.shapes.add_textbox(Inches(8.6), Inches(1.9), Inches(3.9), Inches(4.9))
    tf = tb.text_frame
    tf.word_wrap = True

    bullets = [
        ("Multi-Item Quote Builder", "Add unlimited line items with quantity, unit prices in PKR, and itemized discount percentages."),
        ("Automated Tax Computation", "Dynamic sales tax calculations ensuring 100% pricing accuracy before proposals reach clients."),
        ("Corporate Letterhead", "Pre-formatted with Nexora branding, company contact, NTN, bank details, and customized terms & conditions."),
        ("Instant Quotation-to-Invoice", "Convert approved quotations into live sales orders and invoices with a single click, saving hours of manual data entry.")
    ]
    for i, (title, desc) in enumerate(bullets):
        p_t = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p_t.text = f"• {title}"
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = SKY_BLUE
        p_t.space_after = Pt(2)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_LIGHT
        p_d.space_after = Pt(14)

    # ==========================================
    # SLIDE 5: Multi-Godown Logistics & Warehouses
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, prs)
    add_header(s5, "Supply Chain", "Multi-Godown Storage & Logistics Hubs", "Centralized inventory control across Karachi, Lahore, Faisalabad, and Islamabad facilities.")

    add_image_with_frame(s5, os.path.join(SCREENSHOT_DIR, "06_warehouses.png"), Inches(0.8), Inches(1.9), Inches(7.5), Inches(4.9))

    tb = s5.shapes.add_textbox(Inches(8.6), Inches(1.9), Inches(3.9), Inches(4.9))
    tf = tb.text_frame
    tf.word_wrap = True

    bullets = [
        ("Facility Utilization Tracking", "Monitor capacity utilization percentages for Port Qasim, Lahore Multan Rd, and Islamabad depots."),
        ("Inter-Warehouse Stock Routing", "Execute internal stock transfer requests with digital dispatch notes and receiving acknowledgements."),
        ("Real-Time Valuation in PKR", "Live total inventory valuation tracked by individual warehouse location to eliminate dead stock."),
        ("Facility Manager Governance", "Assign dedicated warehouse supervisors and godown managers with isolated access control.")
    ]
    for i, (title, desc) in enumerate(bullets):
        p_t = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p_t.text = f"• {title}"
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = EMERALD
        p_t.space_after = Pt(2)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_LIGHT
        p_d.space_after = Pt(14)

    # ==========================================
    # SLIDE 6: Product Catalog & Stock Valuation
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, prs)
    add_header(s6, "Inventory Control", "Products Catalog, SKUs & Reorder Alerts", "Comprehensive SKU management, cost vs selling prices, low-stock alerts, and category tagging.")

    add_image_with_frame(s6, os.path.join(SCREENSHOT_DIR, "07_products.png"), Inches(0.8), Inches(1.9), Inches(7.5), Inches(4.9))

    tb = s6.shapes.add_textbox(Inches(8.6), Inches(1.9), Inches(3.9), Inches(4.9))
    tf = tb.text_frame
    tf.word_wrap = True

    bullets = [
        ("Automated Reorder Alerts", "Instant alerts when item quantities drop below safety thresholds, preventing costly stockouts."),
        ("Multi-Price Management", "Track unit purchase costs vs wholesale and retail selling prices with gross margin percentages."),
        ("Categorization & Filtering", "Quickly filter by IT hardware, machinery, raw materials, or consumables with instant barcode / SKU lookup."),
        ("Bulk Import & Export", "Effortlessly import thousands of items from legacy Excel spreadsheets within minutes.")
    ]
    for i, (title, desc) in enumerate(bullets):
        p_t = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p_t.text = f"• {title}"
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = SKY_BLUE
        p_t.space_after = Pt(2)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_LIGHT
        p_d.space_after = Pt(14)

    # ==========================================
    # SLIDE 7: HR, Payroll & EOBI Compliance
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, prs)
    add_header(s7, "Human Capital", "Localized Payroll, EOBI & Direct Salary Credits", "Automated monthly payroll runs with income tax deductions, social security, and 1Link banking.")

    add_image_with_frame(s7, os.path.join(SCREENSHOT_DIR, "08_payroll.png"), Inches(0.8), Inches(1.9), Inches(7.5), Inches(4.9))

    tb = s7.shapes.add_textbox(Inches(8.6), Inches(1.9), Inches(3.9), Inches(4.9))
    tf = tb.text_frame
    tf.word_wrap = True

    bullets = [
        ("Pakistan Tax & EOBI Deductions", "Automated calculations for EOBI employee/employer contributions, PESSI/SESSI, and FBR income tax slabs."),
        ("1-Click Bulk Payroll Generation", "Generate entire company payroll runs across all departments simultaneously with individual adjustments."),
        ("1Link & Raast Direct Bank Transfer", "Export standardized bank salary clearing sheets for Habib Bank (HBL), Meezan, MCB, and UBL."),
        ("Printable Employee Payslips", "Secure digital payslips detailing basic pay, house rent, medical allowance, overtime, and net payout.")
    ]
    for i, (title, desc) in enumerate(bullets):
        p_t = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p_t.text = f"• {title}"
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = EMERALD
        p_t.space_after = Pt(2)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_LIGHT
        p_d.space_after = Pt(14)

    # ==========================================
    # SLIDE 8: Live Attendance & Biometrics
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, prs)
    add_header(s8, "Workforce Monitoring", "Live Attendance & Shift Management", "Real-time biometric radar tracking present, late arrivals, half-days, and leave balances.")

    add_image_with_frame(s8, os.path.join(SCREENSHOT_DIR, "09_attendance.png"), Inches(0.8), Inches(1.9), Inches(7.5), Inches(4.9))

    tb = s8.shapes.add_textbox(Inches(8.6), Inches(1.9), Inches(3.9), Inches(4.9))
    tf = tb.text_frame
    tf.word_wrap = True

    bullets = [
        ("Real-Time Workforce Radar", "Live attendance status counters showing present headcount, late arrivals, on-duty, and absentees."),
        ("Multi-Shift Configurations", "Support for morning, evening, and rotating factory / office shifts with customizable grace periods."),
        ("Overtime & Deduction Rules", "Seamless sync between daily attendance logs and the monthly payroll engine to calculate late deductions."),
        ("Biometric & Mobile Check-In", "Ready for hardware biometric machine integration and mobile field-staff check-in.")
    ]
    for i, (title, desc) in enumerate(bullets):
        p_t = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p_t.text = f"• {title}"
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = SKY_BLUE
        p_t.space_after = Pt(2)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_LIGHT
        p_d.space_after = Pt(14)

    # ==========================================
    # SLIDE 9: Marketing & Ads P&L Analysis
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9, prs)
    add_header(s9, "Commercial Intelligence", "Marketing & Ads Performance Integrated with P&L", "Comprehensive analysis across Meta, Google, and TikTok integrated directly with financials.")

    add_image_with_frame(s9, os.path.join(SCREENSHOT_DIR, "10_marketing.png"), Inches(0.8), Inches(1.9), Inches(7.5), Inches(4.9))

    tb = s9.shapes.add_textbox(Inches(8.6), Inches(1.9), Inches(3.9), Inches(4.9))
    tf = tb.text_frame
    tf.word_wrap = True

    bullets = [
        ("Cross-Platform Attribution", "Track ad spend, impressions, clicks, customer acquisition cost (CAC), and ROAS across Meta, Google Ads & TikTok."),
        ("Direct P&L Statement Integration", "Marketing ad spend automatically reflects as operational expense in financial reports and EBITDA calculations."),
        ("Realized vs Ad-Generated Revenue", "Compare actual cash collections in bank against digital campaign conversions to gauge genuine profitability."),
        ("Channel Efficiency Benchmarking", "Identify which marketing platforms deliver the highest ROI for your products and services.")
    ]
    for i, (title, desc) in enumerate(bullets):
        p_t = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p_t.text = f"• {title}"
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = EMERALD
        p_t.space_after = Pt(2)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_LIGHT
        p_d.space_after = Pt(14)

    # ==========================================
    # SLIDE 10: Security & Single Sign-On
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10, prs)
    add_header(s10, "Enterprise Security", "Granular Access Control & Google Single Sign-On", "Role-based permission architecture protecting your commercial and financial data.")

    add_image_with_frame(s10, os.path.join(SCREENSHOT_DIR, "02_login.png"), Inches(0.8), Inches(1.9), Inches(7.5), Inches(4.9))

    tb = s10.shapes.add_textbox(Inches(8.6), Inches(1.9), Inches(3.9), Inches(4.9))
    tf = tb.text_frame
    tf.word_wrap = True

    bullets = [
        ("Google Workspace SSO", "Secure 1-click single sign-on via corporate Google accounts with enterprise OAuth verification."),
        ("Role-Based Access Control (RBAC)", "Pre-configured roles (Super Admin, HR, Finance, Inventory, Sales, Department Manager) with View/Create/Edit/Approve rights."),
        ("Complete Audit Trail", "Every login, document export, price edit, and approval is recorded with user identity, timestamp, and IP tracking."),
        ("Data Sovereignty & Backups", "Automated cloud backups with strict SSL encryption adhering to State Bank of Pakistan data regulations.")
    ]
    for i, (title, desc) in enumerate(bullets):
        p_t = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p_t.text = f"• {title}"
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = SKY_BLUE
        p_t.space_after = Pt(2)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_LIGHT
        p_d.space_after = Pt(14)

    # ==========================================
    # SLIDE 11: Implementation Roadmap & Next Steps
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11, prs)
    add_header(s11, "Deployment Plan", "Fast-Track Implementation & Commercial Next Steps", "Structured 3-phase rollout ensuring zero disruption to your daily business operations.")

    # 3 Column Process Cards
    col_width = Inches(3.64)
    col_gap = Inches(0.4)
    start_left = Inches(0.8)

    phases = [
        ("PHASE 1 (Week 1)", "System Setup & Data Onboarding", [
            "Company profile & branch setup",
            "Chart of accounts configuration",
            "Import inventory & pricing from Excel",
            "Import employee records & salaries"
        ], EMERALD),
        ("PHASE 2 (Week 2)", "Departmental Staff Training", [
            "Executive & Director dashboard training",
            "Invoicing & accounts staff walkthrough",
            "Warehouse storekeeper training",
            "HR & payroll execution run"
        ], SKY_BLUE),
        ("PHASE 3 (Week 3)", "Live Go-Live & Dedicated Support", [
            "Parallel run with existing records",
            "Final data verification & audit signoff",
            "Full enterprise cutover to Nexora",
            "Dedicated WhatsApp support hotline"
        ], ROYAL_BLUE),
    ]

    for idx, (p_title, p_sub, items, color) in enumerate(phases):
        c_left = start_left + idx * (col_width + col_gap)
        c_shape = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, Inches(2.0), col_width, Inches(4.7))
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = CARD_BG
        c_shape.line.color.rgb = color
        c_shape.line.width = Pt(1.5)

        tb_col = s11.shapes.add_textbox(c_left + Inches(0.25), Inches(2.2), col_width - Inches(0.5), Inches(4.3))
        tf_col = tb_col.text_frame
        tf_col.word_wrap = True

        p0 = tf_col.paragraphs[0]
        p0.text = p_title
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = color
        p0.space_after = Pt(4)

        p1 = tf_col.add_paragraph()
        p1.text = p_sub
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE
        p1.space_after = Pt(16)

        for item in items:
            pi = tf_col.add_paragraph()
            pi.text = f"✓  {item}"
            pi.font.size = Pt(11)
            pi.font.color.rgb = TEXT_LIGHT
            pi.space_after = Pt(10)

    # Save presentation
    prs.save(OUT_PPTX)
    print(f"Presentation successfully created at: {OUT_PPTX}")

if __name__ == "__main__":
    build_presentation()
